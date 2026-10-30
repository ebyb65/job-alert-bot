import asyncio
import json
import os
import re
import time
import urllib.parse
from datetime import datetime, timedelta, timezone

import requests
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches, Pt
from google import genai
from telegram import Bot

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN", "").strip()
TELEGRAM_CHAT_ID = os.getenv("TELEGRAM_CHAT_ID", "").strip()
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash").strip()

PROFILE_FILE = "candidate_profile.json"
HISTORY_FILE = "sent_jobs.json"
MIN_MATCH_SCORE = int(os.getenv("MIN_MATCH_SCORE", "78"))
LOOKBACK_DAYS = int(os.getenv("HISTORY_LOOKBACK_DAYS", "30"))
MAX_JOBS = int(os.getenv("MAX_JOBS", "30"))

client = genai.Client(api_key=GEMINI_API_KEY) if GEMINI_API_KEY else None


def load_json(path, default):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return default


def save_json(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def load_profile():
    return load_json(PROFILE_FILE, {})


def load_sent_jobs():
    data = load_json(HISTORY_FILE, {})
    cutoff = datetime.now(timezone.utc) - timedelta(days=LOOKBACK_DAYS)
    cleaned = {}
    for k, v in data.items():
        try:
            dt = datetime.fromisoformat(v)
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            if dt > cutoff:
                cleaned[k] = v
        except Exception:
            continue
    return cleaned


def save_sent_job(job_id):
    if not job_id:
        return
    sent = load_sent_jobs()
    sent[str(job_id)] = datetime.now(timezone.utc).isoformat()
    save_json(HISTORY_FILE, sent)


def extract_job_id(url):
    m = re.search(r"/(?:view/)?(\d{8,12})(?:/|$|\?)", url or "")
    if not m:
        m = re.search(r"(\d{8,12})", url or "")
    return m.group(1) if m else None


def normalize_text(text):
    return re.sub(r"\s+", " ", (text or "")).strip()


def fetch_job_description(job):
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124 Safari/537.36",
        "Accept-Language": "en-US,en;q=0.9",
    }
    try:
        r = requests.get(job["link"], headers=headers, timeout=15)
        if r.status_code != 200:
            return ""
        soup = BeautifulSoup(r.text, "html.parser")
        node = (
            soup.select_one("div.show-more-less-html__markup")
            or soup.select_one("div.description__text")
            or soup.select_one("section.description")
        )
        return normalize_text(node.get_text(" ", strip=True)) if node else ""
    except Exception:
        return ""


def fetch_linkedin_jobs(profile, total_limit=MAX_JOBS):
    keywords = profile.get("target_roles") or ["HR Operations", "HR Executive", "HR Officer"]
    locations = profile.get("preferred_locations") or ["Eastern Province", "Saudi Arabia"]
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124 Safari/537.36",
        "Accept-Language": "en-US,en;q=0.9",
    }
    sent = load_sent_jobs()
    jobs = []

    # Prefer focused Eastern Province searches first, then Saudi Arabia broadly.
    queries = []
    for loc in locations[:6]:
        for kw in keywords[:8]:
            queries.append((kw, loc))
    queries.extend((kw, "Saudi Arabia") for kw in keywords[:8])

    for kw, loc in queries:
        if len(jobs) >= total_limit:
            break
        url = (
            "https://www.linkedin.com/jobs-guest/jobs/api/seeMoreJobPostings/search"
            f"?keywords={urllib.parse.quote(kw)}"
            f"&location={urllib.parse.quote(loc)}"
            "&f_TPR=r604800&start=0"
        )
        try:
            res = requests.get(url, headers=headers, timeout=15)
            if res.status_code != 200:
                continue
            soup = BeautifulSoup(res.text, "html.parser")
            for card in soup.find_all("li"):
                title = card.find("h3", class_="base-search-card__title")
                company = card.find("h4", class_="base-search-card__subtitle")
                location = card.find("span", class_="job-search-card__location")
                link = card.find("a", class_="base-card__full-link")
                if not title or not link:
                    continue
                title_text = normalize_text(title.get_text())
                lower = title_text.lower()
                if any(x in lower for x in ["tamheer", "internship", "graduate trainee"]):
                    continue
                clean_link = link.get("href", "").split("?")[0]
                job_id = extract_job_id(clean_link)
                if job_id and job_id in sent:
                    continue
                if any(j["link"] == clean_link for j in jobs):
                    continue
                jobs.append({
                    "title": title_text,
                    "company": normalize_text(company.get_text()) if company else "Company Not Listed",
                    "location": normalize_text(location.get_text()) if location else loc,
                    "link": clean_link,
                    "job_id": job_id,
                    "search_keyword": kw,
                })
                if len(jobs) >= total_limit:
                    break
        except Exception as e:
            print(f"Search failed for {kw} / {loc}: {e}")
        time.sleep(0.5)
    return jobs


def heuristic_score(job, profile, description):
    text = f"{job['title']} {job['location']} {description}".lower()
    role_terms = [x.lower() for x in profile.get("target_roles", [])]
    verified = " ".join(profile.get("verified_experience", [])).lower()
    systems = [x.lower() for x in profile.get("systems", [])]

    score = 45
    if any(term in text for term in role_terms):
        score += 15
    for kw in ["hr operations", "onboarding", "payroll", "employee records", "personnel", "medical insurance", "recruitment", "qiwa", "muqeem", "gosi", "employee relations", "workforce"]:
        if kw in text and kw in verified:
            score += 3
    if any(sys in text for sys in systems):
        score += 5
    if any(loc.lower() in job["location"].lower() for loc in profile.get("preferred_locations", [])):
        score += 8
    if any(x in text for x in ["saudi nationals only", "saudi only", "saudi national only"]):
        score = 0
    if any(x in text for x in ["fresh graduate", "internship", "tamheer"]):
        score = min(score, 20)
    return max(0, min(100, score))


def evaluate_with_ai(job, profile, description):
    if not client:
        return None
    prompt = f"""
You are an ATS/recruitment analyst. Evaluate the vacancy strictly against the VERIFIED candidate profile below.
Never invent experience, qualifications, certifications, industries, countries, languages, availability, salary, or job duties.
If the vacancy requests something not in the profile, list it as a gap instead of implying the candidate has it.

CANDIDATE PROFILE (verified):
{json.dumps(profile, ensure_ascii=False, indent=2)}

VACANCY:
Title: {job['title']}
Company: {job['company']}
Location: {job['location']}
Description: {description or 'Description could not be retrieved; evaluate conservatively from title/location only.'}

Return JSON only with:
{{
  "should_apply": true,
  "match_score": 0,
  "match_reason": "short factual reason",
  "gaps": ["gap 1"],
  "ats_keywords": ["keyword 1", "keyword 2", "keyword 3", "keyword 4", "keyword 5", "keyword 6"],
  "ats_summary": "2-3 sentence summary using only verified facts",
  "tailored_bullets": ["4-6 truthful bullets, each supported by the verified profile"],
  "cover_letter": "short natural 2-paragraph cover letter using only verified facts"
}}
Scoring guidance: 90-100 exceptional direct match; 80-89 strong; 70-79 reasonable; below 70 weak. Penalize mandatory degree/certification/nationality/location requirements that are not met or unknown.
"""
    try:
        response = client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt,
            config={"response_mime_type": "application/json"},
        )
        return json.loads(response.text)
    except Exception as e:
        print(f"AI evaluation failed: {e}")
        return None


def generate_complete_word_cv(job_title, company, profile, data, filename):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.65)
        s.right_margin = Inches(0.65)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)

    p = doc.add_paragraph()
    p.add_run(profile.get("name", "Eby Benjamin").upper()).bold = True
    p.runs[0].font.size = Pt(18)
    p = doc.add_paragraph()
    p.add_run(f"SENIOR HR OPERATIONS & WORKFORCE ADMINISTRATION | TARGET: {job_title}").bold = True
    p = doc.add_paragraph(
        f"{profile.get('location','')} | {profile.get('mobile','')} | {profile.get('email','')} | "
        + " | ".join(profile.get("work_authorization", []))
    )

    def h(text):
        q = doc.add_paragraph()
        q.add_run(text).bold = True
        q.runs[0].font.size = Pt(10.5)

    h("PROFESSIONAL SUMMARY")
    doc.add_paragraph(data.get("ats_summary") or profile.get("summary", ""))

    h("CORE COMPETENCIES")
    doc.add_paragraph(" • ".join(data.get("ats_keywords", [])))
    doc.add_paragraph("Systems: " + ", ".join(profile.get("systems", [])) + " | Tools: " + ", ".join(profile.get("tools", [])))

    h("SELECTED RELEVANT EXPERIENCE")
    for bullet in data.get("tailored_bullets", [])[:6]:
        doc.add_paragraph(bullet, style="List Bullet")

    h("VERIFIED CAREER HIGHLIGHTS")
    for bullet in profile.get("verified_experience", [])[:10]:
        doc.add_paragraph(bullet, style="List Bullet")

    h("EDUCATION")
    for item in profile.get("education", []):
        doc.add_paragraph(item)

    h("LANGUAGES")
    doc.add_paragraph(" | ".join(f"{k}: {v}" for k, v in profile.get("languages", {}).items()))

    doc.save(filename)
    return filename


async def notify(bot, text):
    await bot.send_message(chat_id=TELEGRAM_CHAT_ID, text=text, disable_web_page_preview=True)


async def run_live_agent():
    if not TELEGRAM_BOT_TOKEN or not TELEGRAM_CHAT_ID:
        raise RuntimeError("TELEGRAM_BOT_TOKEN and TELEGRAM_CHAT_ID are required")

    profile = load_profile()
    if not profile:
        raise RuntimeError("candidate_profile.json is missing or invalid")

    bot = Bot(token=TELEGRAM_BOT_TOKEN)
    jobs = fetch_linkedin_jobs(profile)
    print(f"Discovered {len(jobs)} new vacancies")

    strong_matches = []
    for idx, job in enumerate(jobs, start=1):
        print(f"Processing {idx}/{len(jobs)}: {job['title']} at {job['company']}")
        description = fetch_job_description(job)
        heuristic = heuristic_score(job, profile, description)
        if heuristic < 60:
            print(f"Filtered before AI: heuristic {heuristic}")
            if job.get("job_id"):
                save_sent_job(job["job_id"])
            continue

        data = evaluate_with_ai(job, profile, description)
        if data is None:
            data = {
                "should_apply": heuristic >= MIN_MATCH_SCORE,
                "match_score": heuristic,
                "match_reason": "AI unavailable; score based on title, location and verified profile keyword overlap.",
                "gaps": [],
                "ats_keywords": [],
                "ats_summary": profile.get("summary", ""),
                "tailored_bullets": profile.get("verified_experience", [])[:5],
                "cover_letter": "",
            }

        score = int(data.get("match_score", 0))
        if data.get("should_apply") and score >= MIN_MATCH_SCORE:
            strong_matches.append((score, job, data))
        elif job.get("job_id"):
            save_sent_job(job["job_id"])
        await asyncio.sleep(2)

    strong_matches.sort(key=lambda x: x[0], reverse=True)

    if not strong_matches:
        await notify(bot, "No new high-confidence HR job matches were found today.")
        return

    for score, job, data in strong_matches[:8]:
        app_url = f"https://www.linkedin.com/jobs/view/{job['job_id']}/" if job.get("job_id") else job["link"]
        gaps = "; ".join(data.get("gaps", [])[:4]) or "No major gap identified from the available job description."
        keywords = ", ".join(data.get("ats_keywords", [])[:6]) or "Not extracted"
        msg = (
            f"🎯 {score}% MATCH\n"
            f"{job['title']} | {job['company']}\n"
            f"📍 {job['location']}\n"
            f"Why: {data.get('match_reason','Strong overlap with verified experience.')}\n"
            f"Gaps: {gaps}\n"
            f"ATS: {keywords}\n"
            f"Apply: {app_url}"
        )
        await notify(bot, msg)

        safe_co = re.sub(r"[^A-Za-z0-9_-]+", "_", job["company"])[:24].strip("_") or "Company"
        filename = f"Eby_Benjamin_{safe_co}_{score}.docx"
        generate_complete_word_cv(job["title"], job["company"], profile, data, filename)
        with open(filename, "rb") as f:
            await bot.send_document(
                chat_id=TELEGRAM_CHAT_ID,
                document=f,
                caption=f"ATS-tailored CV | {job['title']} | {job['company']} | {score}% match",
            )
        os.remove(filename)

        cover = data.get("cover_letter", "").strip()
        if cover:
            await notify(bot, "Cover letter:\n\n" + cover)
        if job.get("job_id"):
            save_sent_job(job["job_id"])
        await asyncio.sleep(2)


if __name__ == "__main__":
    asyncio.run(run_live_agent())
